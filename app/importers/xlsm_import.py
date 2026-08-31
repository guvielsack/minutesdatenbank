from datetime import date, datetime, time
from io import BytesIO
from typing import BinaryIO

import openpyxl
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from app.config import DEFAULT_PROJECT
from app.models import ImportLog, MinuteEntry, ProjectSettings, YearPlanEntry
from app.services.calendar import excel_serial_to_date, schulferien_label
from app.services.minutes_actions import _normalize_type

MINUTES_SHEET = "Minutes"
YEAR_PLAN_SHEET_NAMES = ("Jahresdetail-Plan", "Jahresdetail-Planung")


def _cell_date(value) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return excel_serial_to_date(value)
    return None


def _project_settings_from_workbook(wb: Workbook) -> dict[str, str]:
    ws = wb[MINUTES_SHEET]
    name = ws["C1"].value or DEFAULT_PROJECT["name"]
    if isinstance(name, str):
        name = name.strip().strip("'")

    return {
        "name": str(name),
        "team_label": str(ws["D1"].value or DEFAULT_PROJECT["team_label"]),
        "meeting": str(ws["E1"].value or DEFAULT_PROJECT["meeting"]),
        "participants": str(ws["F1"].value or DEFAULT_PROJECT["participants"]),
    }


def _ensure_project(db: Session, wb: Workbook) -> ProjectSettings:
    values = _project_settings_from_workbook(wb)
    settings = db.query(ProjectSettings).first()
    if settings:
        for field, value in values.items():
            setattr(settings, field, value)
        db.commit()
        db.refresh(settings)
        return settings

    settings = ProjectSettings(**values)
    db.add(settings)
    db.commit()
    db.refresh(settings)
    return settings


def _find_year_plan_sheet(wb: Workbook) -> Worksheet | None:
    for sheet_name in YEAR_PLAN_SHEET_NAMES:
        if sheet_name in wb.sheetnames:
            return wb[sheet_name]
    return None


def _require_minutes_sheet(wb: Workbook) -> None:
    if MINUTES_SHEET not in wb.sheetnames:
        raise ValueError(
            f'Sheet "{MINUTES_SHEET}" fehlt in der Datei. '
            f'Erwartet: "{MINUTES_SHEET}" (Pflicht), optional einer von: '
            f'"{YEAR_PLAN_SHEET_NAMES[0]}" oder "{YEAR_PLAN_SHEET_NAMES[1]}". '
            f"Vorhandene Sheets: {', '.join(wb.sheetnames) or '(keine)'}"
        )


def import_minutes_sheet(db: Session, wb: Workbook) -> int:
    ws = wb[MINUTES_SHEET]
    db.query(MinuteEntry).delete()
    count = 0

    for row_idx in range(4, ws.max_row + 1):
        entry_type = _normalize_type(ws.cell(row_idx, 2).value)
        content = ws.cell(row_idx, 3).value
        if not entry_type and not content:
            continue

        entry = MinuteEntry(
            row_nr=int(ws.cell(row_idx, 1).value or count),
            entry_type=entry_type or "A",
            content=str(content) if content is not None else None,
            responsible=_as_str(ws.cell(row_idx, 4).value),
            along_with=_as_str(ws.cell(row_idx, 5).value),
            since_when=_cell_date(ws.cell(row_idx, 6).value),
            until_when=_cell_date(ws.cell(row_idx, 7).value),
            remarks=_cell_remarks_or_duration(ws.cell(row_idx, 8).value),
            category=_as_str(ws.cell(row_idx, 9).value),
            status=_as_str(ws.cell(row_idx, 10).value),
            link=_as_str(ws.cell(row_idx, 11).value),
        )
        db.add(entry)
        count += 1

    db.commit()
    return count


def import_year_plan_sheet(db: Session, ws: Worksheet) -> int:
    db.query(YearPlanEntry).delete()
    count = 0

    for row_idx in range(3, ws.max_row + 1):
        meeting_date = _resolve_plan_date(ws, row_idx)
        focus = ws.cell(row_idx, 4).value
        if meeting_date is None and not focus:
            continue

        entry = YearPlanEntry(
            row_nr=count,
            weekday=str(ws.cell(row_idx, 1).value or "Dienstag"),
            meeting_date=meeting_date or date.today(),
            focus_topic=_as_str(focus),
            fach_topic=_as_str(ws.cell(row_idx, 5).value),
            location=_as_str(ws.cell(row_idx, 6).value),
            absence_ao=_absence(ws.cell(row_idx, 7).value),
            absence_cn=_absence(ws.cell(row_idx, 8).value),
            absence_dk=_absence(ws.cell(row_idx, 9).value),
            absence_js=_absence(ws.cell(row_idx, 10).value),
            absence_gv=_absence(ws.cell(row_idx, 11).value),
            absence_jb=_absence(ws.cell(row_idx, 12).value),
            absence_sm=_absence(ws.cell(row_idx, 13).value),
            absence_re=_absence(ws.cell(row_idx, 14).value),
            absence_ak=_absence(ws.cell(row_idx, 15).value),
        )
        db.add(entry)
        count += 1

    db.commit()
    return count


def _resolve_plan_date(ws, row_idx: int) -> date | None:
    value = ws.cell(row_idx, 2).value
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _cell_remarks_or_duration(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, (int, float)):
        total_minutes = int(round(float(value) * 24 * 60))
        hours, minutes = divmod(total_minutes, 60)
        return f"{hours:02d}:{minutes:02d}"
    text = str(value).strip()
    return text or None


def _as_str(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _absence(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if text.upper() == "X":
        return "X"
    return None


def _load_workbook(source: str | BinaryIO) -> Workbook:
    return openpyxl.load_workbook(source, data_only=True, keep_vba=True)


def _import_openpyxl_workbook(db: Session, wb: Workbook, source_file: str) -> tuple[ImportLog, bool]:
    _require_minutes_sheet(wb)
    _ensure_project(db, wb)
    minute_rows = import_minutes_sheet(db, wb)

    year_plan_ws = _find_year_plan_sheet(wb)
    year_plan_included = year_plan_ws is not None
    year_plan_rows = import_year_plan_sheet(db, year_plan_ws) if year_plan_included else 0

    log = ImportLog(
        source_file=source_file,
        minute_rows=minute_rows,
        year_plan_rows=year_plan_rows,
    )
    db.add(log)
    db.commit()
    db.refresh(log)
    return log, year_plan_included


def import_workbook(db: Session, path: str) -> tuple[ImportLog, bool]:
    wb = _load_workbook(path)
    return _import_openpyxl_workbook(db, wb, path)


def import_workbook_bytes(db: Session, content: bytes, filename: str) -> tuple[ImportLog, bool]:
    wb = _load_workbook(BytesIO(content))
    return _import_openpyxl_workbook(db, wb, filename)


def resolve_school_holiday(entry: YearPlanEntry) -> str | None:
    if entry.school_holiday:
        return entry.school_holiday
    return schulferien_label(entry.meeting_date)


def year_plan_with_holidays(entry: YearPlanEntry) -> dict:
    data = {column.name: getattr(entry, column.name) for column in entry.__table__.columns}
    data["school_holiday"] = resolve_school_holiday(entry)
    return data
